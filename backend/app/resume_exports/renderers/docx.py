from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.resume_exports.renderers.base import ResumeExportRenderer
from app.resume_exports.renderers.models import RenderedResume
from app.resumes.models import Resume


class DocxResumeRenderer(ResumeExportRenderer):
    """
    Render a complete resume as a DOCX document.

    The renderer receives a fully loaded Resume entity and converts
    it into a structured Microsoft Word document.
    """

    @property
    def media_type(self) -> str:
        return (
            "application/"
            "vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    @property
    def file_extension(self) -> str:
        return "docx"

    def render(
        self,
        resume: Resume,
    ) -> RenderedResume:
        document = Document()

        self._configure_document(
            document,
        )

        self._render_header(
            document,
            resume,
        )

        self._render_summary(
            document,
            resume,
        )

        self._render_experience(
            document,
            resume,
        )

        self._render_education(
            document,
            resume,
        )

        self._render_skills(
            document,
            resume,
        )

        self._render_projects(
            document,
            resume,
        )

        self._render_certifications(
            document,
            resume,
        )

        self._render_generated_content(
            document,
            resume,
        )

        document_buffer = BytesIO()

        document.save(
            document_buffer,
        )

        return RenderedResume(
            content=document_buffer.getvalue(),
            filename=self._build_filename(resume),
            media_type=self.media_type,
        )

    def _configure_document(
        self,
        document: Document,
    ) -> None:
        section = document.sections[0]

        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

        normal_style = document.styles["Normal"]

        normal_style.font.name = "Arial"
        normal_style.font.size = Pt(10)

    def _render_header(
        self,
        document: Document,
        resume: Resume,
    ) -> None:
        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            resume.title.strip()
            if resume.title and resume.title.strip()
            else "Resume",
        )

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(20)

    def _render_summary(
        self,
        document: Document,
        resume: Resume,
    ) -> None:
        if not resume.summary:
            return

        self._add_section_heading(
            document,
            "Professional Summary",
        )

        paragraph = document.add_paragraph(
            resume.summary,
        )

        paragraph.paragraph_format.space_after = Pt(8)

    def _render_experience(
        self,
        document: Document,
        resume: Resume,
    ) -> None:
        if not resume.experiences:
            return

        self._add_section_heading(
            document,
            "Experience",
        )

        for experience in resume.experiences:
            paragraph = document.add_paragraph()

            job_title = paragraph.add_run(
                experience.job_title,
            )

            job_title.bold = True

            company = paragraph.add_run(
                f" | {experience.company}",
            )

            company.bold = True

            if experience.location:
                paragraph.add_run(
                    f" | {experience.location}",
                )

            date_paragraph = document.add_paragraph()

            date_paragraph.paragraph_format.space_after = Pt(2)

            date_paragraph.add_run(
                self._format_date_range(
                    experience.start_date,
                    experience.end_date,
                    experience.is_current,
                ),
            )

            if experience.description:
                description = document.add_paragraph(
                    experience.description,
                )

                description.paragraph_format.space_after = Pt(6)

    def _render_education(
        self,
        document: Document,
        resume: Resume,
    ) -> None:
        if not resume.educations:
            return

        self._add_section_heading(
            document,
            "Education",
        )

        for education in resume.educations:
            paragraph = document.add_paragraph()

            degree = paragraph.add_run(
                education.degree,
            )

            degree.bold = True

            if education.field_of_study:
                paragraph.add_run(
                    f" in {education.field_of_study}",
                )

            institution = document.add_paragraph()

            institution.add_run(
                education.institution,
            ).bold = True

            if education.location:
                institution.add_run(
                    f" | {education.location}",
                )

            date_paragraph = document.add_paragraph()

            date_paragraph.paragraph_format.space_after = Pt(2)

            date_paragraph.add_run(
                self._format_date_range(
                    education.start_date,
                    education.end_date,
                    education.is_current,
                ),
            )

            if education.grade:
                grade_paragraph = document.add_paragraph(
                    f"Grade: {education.grade}",
                )

                grade_paragraph.paragraph_format.space_after = Pt(2)

            if education.description:
                description = document.add_paragraph(
                    education.description,
                )

                description.paragraph_format.space_after = Pt(6)

    def _render_skills(
        self,
        document: Document,
        resume: Resume,
    ) -> None:
        if not resume.skills:
            return

        self._add_section_heading(
            document,
            "Skills",
        )

        skills = [
            skill.name
            for skill in resume.skills
            if skill.name
        ]

        if not skills:
            return

        paragraph = document.add_paragraph(
            " • ".join(skills),
        )

        paragraph.paragraph_format.space_after = Pt(8)

    def _render_projects(
        self,
        document: Document,
        resume: Resume,
    ) -> None:
        if not resume.projects:
            return

        self._add_section_heading(
            document,
            "Projects",
        )

        for project in resume.projects:
            paragraph = document.add_paragraph()

            name = paragraph.add_run(
                project.name,
            )

            name.bold = True

            if project.technologies:
                technologies = document.add_paragraph(
                    f"Technologies: {project.technologies}",
                )

                technologies.paragraph_format.space_after = Pt(2)

            if project.description:
                description = document.add_paragraph(
                    project.description,
                )

                description.paragraph_format.space_after = Pt(2)

            if project.project_url:
                document.add_paragraph(
                    f"Project: {project.project_url}",
                )

            if project.repository_url:
                document.add_paragraph(
                    f"Repository: {project.repository_url}",
                )

            date_paragraph = document.add_paragraph()

            date_paragraph.paragraph_format.space_after = Pt(6)

            date_paragraph.add_run(
                self._format_date_range(
                    project.start_date,
                    project.end_date,
                    project.is_ongoing,
                ),
            )

    def _render_certifications(
        self,
        document: Document,
        resume: Resume,
    ) -> None:
        if not resume.certifications:
            return

        self._add_section_heading(
            document,
            "Certifications",
        )

        for certification in resume.certifications:
            paragraph = document.add_paragraph()

            name = paragraph.add_run(
                certification.name,
            )

            name.bold = True

            if certification.issuing_organization:
                paragraph.add_run(
                    f" | {certification.issuing_organization}",
                )

            if certification.credential_id:
                document.add_paragraph(
                    f"Credential ID: {certification.credential_id}",
                )

            if certification.credential_url:
                document.add_paragraph(
                    f"Credential URL: {certification.credential_url}",
                )

            date_paragraph = document.add_paragraph()

            date_paragraph.paragraph_format.space_after = Pt(6)

            date_paragraph.add_run(
                self._format_certification_dates(
                    issue_date=certification.issue_date,
                    expiration_date=certification.expiration_date,
                    does_not_expire=certification.does_not_expire,
                ),
            )

    def _render_generated_content(
        self,
        document: Document,
        resume: Resume,
    ) -> None:
        if not resume.generated_content:
            return

        self._add_section_heading(
            document,
            "Generated Resume Content",
        )

        document.add_paragraph(
            resume.generated_content,
        )

    def _add_section_heading(
        self,
        document: Document,
        title: str,
    ) -> None:
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)

        run = paragraph.add_run(
            title,
        )

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(13)

    @staticmethod
    def _format_date_range(
        start_date,
        end_date,
        is_current: bool,
    ) -> str:
        start = start_date.strftime("%B %Y")

        if is_current:
            return f"{start} - Present"

        if end_date is None:
            return start

        return f"{start} - {end_date.strftime('%B %Y')}"

    @staticmethod
    def _format_certification_dates(
        *,
        issue_date,
        expiration_date,
        does_not_expire: bool,
    ) -> str:
        issued = f"Issued: {issue_date.strftime('%B %Y')}"

        if does_not_expire:
            return f"{issued} | Does not expire"

        if expiration_date is None:
            return issued

        return (
            f"{issued} | "
            f"Expires: {expiration_date.strftime('%B %Y')}"
        )

    def _build_filename(
        self,
        resume: Resume,
    ) -> str:
        filename = resume.title.strip()

        if not filename:
            filename = "resume"

        sanitized = "".join(
            character
            if character.isalnum() or character in {" ", "-", "_"}
            else "_"
            for character in filename
        )

        sanitized = "_".join(
            sanitized.split(),
        )

        return f"{sanitized}.docx"