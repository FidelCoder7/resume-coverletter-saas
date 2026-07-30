from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document

from app.resume_exports.renderers import DocxResumeRenderer
from tests.factories.resume_factory import make_resume


@pytest.fixture
def resume():
    resume = make_resume(
    user_id="00000000-0000-0000-0000-000000000001",
    title="Backend Resume",
    summary="Python Backend Developer",
    )


    resume.experiences = []
    resume.educations = []
    resume.skills = []
    resume.projects = []
    resume.certifications = []

    return resume


def test_docx_renderer_metadata():
    renderer = DocxResumeRenderer()


    assert (
        renderer.media_type
        == "application/"
        "vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert renderer.file_extension == "docx"


def test_docx_renderer_produces_valid_docx(
    resume,
):
    renderer = DocxResumeRenderer()


    result = renderer.render(
        resume,
    )

    assert result.content
    assert result.filename.endswith(".docx")
    assert result.media_type == renderer.media_type

    document = Document(
        BytesIO(result.content),
    )

    assert document.paragraphs


def test_docx_renderer_contains_resume_content(
    resume,
):
    renderer = DocxResumeRenderer()


    result = renderer.render(
        resume,
    )

    document = Document(
        BytesIO(result.content),
    )

    document_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert resume.title in document_text

    if resume.summary:
        assert resume.summary in document_text

    for experience in resume.experiences:
        assert experience.company in document_text
        assert experience.job_title in document_text

    for education in resume.educations:
        assert education.institution in document_text

    for skill in resume.skills:
        assert skill.name in document_text

    for project in resume.projects:
        assert project.name in document_text

    for certification in resume.certifications:
        assert certification.name in document_text


def test_docx_renderer_output_is_zip_archive(
    resume,
):
    renderer = DocxResumeRenderer()


    result = renderer.render(
        resume,
    )

    with ZipFile(
        BytesIO(result.content),
    ) as archive:
        names = archive.namelist()

    assert "[Content_Types].xml" in names
    assert "word/document.xml" in names


def test_docx_renderer_uses_fallback_filename_for_empty_title(
    resume,
):
    resume.title = ""


    renderer = DocxResumeRenderer()

    result = renderer.render(
        resume,
    )

    assert result.filename == "resume.docx"
